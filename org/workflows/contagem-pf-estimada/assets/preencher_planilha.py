#!/usr/bin/env python3
"""Preenche o template da planilha de contagem de PF a partir de um JSON.

Uso: python3 preencher_planilha.py contagem.json saida.xlsx

JSON:
{
  "os": "OS.2026.294", "aplicacao": "Obrasim", "escopo": "texto do escopo",
  "tipo_contagem": "Projeto de Melhoria", "nivel": "Estimativa (NESMA)",
  "responsavel": "Nome", "tecnologia": "", "data": "2026-09-24",
  "documentacao": ["HU06.06 - ..."],
  "funcoes": [
    {"nome": "Listar contratos", "tipo": "SE", "manutencao": "A50",
     "td": null, "tr": null, "referencia": "", "observacao": "HU06.06 - ..."}
  ]
}

Fórmulas do template ficam intactas (o Excel recalcula ao abrir). O script
replica o cálculo só para validar a entrada e imprimir o resumo.
"""
import datetime
import json
import sys
from pathlib import Path

import openpyxl

TEMPLATE = Path(__file__).with_name("template-contagem-pf.xlsx")
TIPOS = {"ALI", "AIE", "EE", "SE", "CE"}
TIPOS_CONTAGEM = {"Projeto de Desenvolvimento", "Projeto de Melhoria", "Aplicação"}
PESOS = {"ALI": (7, 10, 15), "AIE": (5, 7, 10), "EE": (3, 4, 6), "SE": (4, 5, 7), "CE": (3, 4, 6)}
PRIMEIRA_LINHA, ULTIMA_LINHA = 8, 522


def complexidade(tipo, td, tr):
    """Mesma regra da coluna I do template (IFPUG; sem TD/TR = NESMA estimada)."""
    if td is None or tr is None:
        return "L" if tipo in ("ALI", "AIE") else "A"
    if tipo == "EE":
        if tr >= 3:
            return "H" if td >= 5 else "A"
        if tr >= 2:
            return "H" if td >= 16 else ("L" if td <= 4 else "A")
        return "L" if td <= 15 else "A"
    if tipo in ("SE", "CE"):
        if tr >= 4:
            return "H" if td >= 6 else "A"
        if tr >= 2:
            return "H" if td >= 20 else ("L" if td <= 5 else "A")
        return "L" if td <= 19 else "A"
    if tr >= 6:
        return "H" if td >= 20 else "A"
    if tr >= 2:
        return "H" if td >= 51 else ("L" if td <= 19 else "A")
    return "L" if td <= 50 else "A"


def ler_deflatores(wb):
    ws = wb["Deflatores"]
    manut, inm = {}, {}
    for r in range(4, 39):
        sigla = ws[f"G{r}"].value
        if sigla and sigla.strip() != ".":
            manut[sigla.strip()] = (ws[f"H{r}"].value or 0, ws[f"I{r}"].value or 0, ws[f"B{r}"].value)
    for r in range(42, 65):
        sigla = ws[f"G{r}"].value
        if sigla and sigla.strip() != ".":
            inm[sigla.strip()] = (ws[f"H{r}"].value or 0, ws[f"B{r}"].value)
    return manut, inm


def main(entrada, saida):
    dados = json.loads(Path(entrada).read_text(encoding="utf-8"))
    wb = openpyxl.load_workbook(TEMPLATE)
    manut, inm = ler_deflatores(wb)

    erros = []
    if dados.get("tipo_contagem") not in TIPOS_CONTAGEM:
        erros.append(f"tipo_contagem inválido: {dados.get('tipo_contagem')!r} (use {sorted(TIPOS_CONTAGEM)})")
    funcoes = dados.get("funcoes", [])
    if not funcoes:
        erros.append("nenhuma função informada")
    if len(funcoes) > ULTIMA_LINHA - PRIMEIRA_LINHA + 1:
        erros.append("funções demais para o template")

    total_ifpug = total_fs = 0
    linhas = []
    for i, f in enumerate(funcoes, 1):
        tipo, sigla = f.get("tipo"), f.get("manutencao")
        if tipo in inm:
            pf, fs, cx = 0, inm[tipo][0], "-"
        elif tipo in TIPOS:
            if sigla not in manut:
                erros.append(f"função {i} ({f.get('nome')}): manutenção {sigla!r} não existe em Deflatores")
                continue
            cx = complexidade(tipo, f.get("td"), f.get("tr"))
            pf = PESOS[tipo]["LAH".index(cx)]
            pct, fixa, _ = manut[sigla]
            fs = pct * pf + fixa
        else:
            erros.append(f"função {i} ({f.get('nome')}): tipo {tipo!r} inválido")
            continue
        total_ifpug += pf
        total_fs += fs
        linhas.append((f.get("nome"), tipo, sigla or "", cx, pf, fs))

    if erros:
        sys.exit("ERRO:\n- " + "\n- ".join(erros))

    identificador = f"{dados['os']} - {dados['aplicacao']}"
    c = wb["Contagem"]
    c["F6"] = dados["tipo_contagem"]
    c["F7"] = dados.get("nivel", "Estimativa (NESMA)")
    c["F9"] = dados.get("responsavel", "")
    c["R7"] = dados.get("tecnologia") or None
    c["R9"] = datetime.date.fromisoformat(dados.get("data") or datetime.date.today().isoformat())
    c["A17"] = dados.get("escopo") or identificador
    c["A22"] = "\n".join(dados.get("documentacao", [])) or None

    ws = wb["Funções"]
    for r, f in enumerate(funcoes, PRIMEIRA_LINHA):
        ws[f"A{r}"] = f["nome"]
        ws[f"B{r}"] = f["tipo"]
        ws[f"C{r}"] = f.get("manutencao") or None
        ws[f"D{r}"] = f.get("td")
        ws[f"E{r}"] = f.get("tr")
        ws[f"N{r}"] = f.get("referencia") or None
        ws[f"O{r}"] = f.get("observacao") or None

    wb.calculation.fullCalcOnLoad = True
    wb.save(saida)

    print(f"{'Função':50} {'Tipo':5} {'Man.':6} {'Cx':3} {'PF':>4} {'FS':>6}")
    for nome, tipo, sigla, cx, pf, fs in linhas:
        print(f"{nome[:50]:50} {tipo:5} {sigla:6} {cx:3} {pf:>4} {fs:>6g}")
    print(f"\nPF IFPUG: {total_ifpug}  |  PF Local da FS: {total_fs:g}")
    print(f"Gerado: {saida}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
